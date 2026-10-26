"""Unified output reporting, backfill, and cleanup helpers."""

from __future__ import annotations

import copy
import ast
import gzip
import json
import math
import os
import pickle
import re
import shutil
from collections.abc import Iterable, Sequence
from dataclasses import asdict, is_dataclass
from pathlib import Path
from typing import Any

try:  # Legacy benchmark workbook support; new runtime workbooks use artifact-tool.
    from openpyxl import Workbook, load_workbook
except ImportError:  # Keep runtime/reporting usable in the minimal study environment.
    Workbook = None
    load_workbook = None

from guiagent.common import sanitize_name
from study.bootstrap import PROJECT_ROOT, ensure_import_paths
from study.localization import (
    EXCEL_HEADERS_ZH,
    ChineseTranslator,
    localize_outputs_tree,
    localize_record,
    localize_run_output_files,
)
from study.runtime_metrics import (
    aggregate_usage,
    capture_state_artifacts,
    normalize_usage,
    reconcile_token_layers,
)


ensure_import_paths()

try:  # AndroidLab-only dependencies are optional for AndroidWorld runs.
    from evaluation.configs import AppConfig
    from evaluation.task import dump_xml
except ImportError:
    AppConfig = None
    dump_xml = None


EXCEL_COLUMNS = [
    "record_id",
    "benchmark",
    "mode",
    "mode_label",
    "run_name",
    "task_id",
    "task",
    "execution_steps",
    "final_result",
    "success",
    "step_count",
    "prompt_tokens",
    "completion_tokens",
    "total_tokens",
    "runtime_seconds",
    "avg_step_seconds",
    "started_at",
    "finished_at",
    "task_report_dir",
    "run_output_dir",
    "notes",
]


def _natural_key(value: str) -> list[Any]:
    return [int(part) if part.isdigit() else part.lower() for part in re.split(r"(\d+)", value)]


def _to_jsonable(obj: Any) -> Any:
    if is_dataclass(obj):
        return _to_jsonable(asdict(obj))
    if isinstance(obj, Path):
        return str(obj)
    if isinstance(obj, dict):
        return {str(key): _to_jsonable(value) for key, value in obj.items()}
    if isinstance(obj, (list, tuple, set)):
        return [_to_jsonable(value) for value in obj]
    if isinstance(obj, (str, int, float, bool)) or obj is None:
        return obj
    if hasattr(obj, "value"):
        return getattr(obj, "value")
    return str(obj)


def _load_json(path: Path) -> dict[str, Any] | None:
    if not path.exists():
        return None
    return json.loads(path.read_text(encoding="utf-8"))


def _write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(_to_jsonable(payload), ensure_ascii=False, indent=2), encoding="utf-8")


def _copy_file(src: Path | None, dst: Path) -> Path | None:
    if src is None or not src.exists():
        return None
    dst.parent.mkdir(parents=True, exist_ok=True)
    try:
        shutil.copy2(src, dst)
    except FileNotFoundError:
        return None
    return dst


def _copy_tree_contents(src_dir: Path | None, dst_dir: Path, *, patterns: Sequence[str]) -> list[str]:
    copied: list[str] = []
    if src_dir is None or not src_dir.exists():
        return copied
    dst_dir.mkdir(parents=True, exist_ok=True)
    for pattern in patterns:
        for src in sorted(src_dir.glob(pattern), key=lambda item: _natural_key(item.name)):
            if not src.is_file():
                continue
            target = dst_dir / src.name
            try:
                shutil.copy2(src, target)
            except FileNotFoundError:
                continue
            copied.append(str(target))
    return copied


def _copy_tree_contents_preserve(src_dir: Path | None, dst_dir: Path, *, patterns: Sequence[str]) -> list[str]:
    copied: list[str] = []
    if src_dir is None or not src_dir.exists():
        return copied
    for pattern in patterns:
        for src in sorted(src_dir.glob(pattern), key=lambda item: _natural_key(str(item))):
            if not src.is_file():
                continue
            try:
                relative = src.relative_to(src_dir)
            except ValueError:
                relative = Path(src.name)
            target = dst_dir / relative
            target.parent.mkdir(parents=True, exist_ok=True)
            try:
                shutil.copy2(src, target)
            except FileNotFoundError:
                continue
            copied.append(str(target))
    return copied


def _safe_mtime(path: Path) -> float:
    try:
        return path.stat().st_mtime
    except FileNotFoundError:
        return 0.0


def _path_from_value(value: Any) -> Path | None:
    if not value or not isinstance(value, (str, os.PathLike)):
        return None
    text = str(value)
    if text.startswith(("<inline_image>", "data:image/", "file://")):
        if text.startswith("file://"):
            return Path(text[7:])
        return None
    # Payloads also contain inline XML/JSON/text.  On macOS, constructing a
    # Path from a long inline XML string can raise ``OSError: [Errno 63]
    # File name too long`` before we get a chance to check ``exists()``.  A
    # non-path value must remain a value in the report, not abort report
    # finalization.
    try:
        path = Path(text)
    except (OSError, ValueError):
        return None
    try:
        return path if path.exists() and path.is_file() else None
    except (OSError, ValueError):
        # ``stat`` can also reject an inline value (for example a long XML
        # document) even after ``Path`` construction succeeds.
        return None


def _copy_step_artifact(src_value: Any, dst_dir: Path, *, prefix: str = "") -> str | None:
    src = _path_from_value(src_value)
    if src is None:
        return None
    dst_dir.mkdir(parents=True, exist_ok=True)
    safe_name = sanitize_name(src.name) or "artifact"
    target = dst_dir / f"{prefix}{safe_name}"
    stem = target.stem
    suffix = target.suffix
    counter = 2
    while target.exists():
        target = dst_dir / f"{stem}_{counter}{suffix}"
        counter += 1
    try:
        shutil.copy2(src, target)
    except FileNotFoundError:
        return None
    return str(target)


def _path_rewrite_keys(src_value: Any, src_path: Path) -> set[str]:
    keys = {str(src_value), str(src_path)}
    try:
        resolved = src_path.resolve()
        keys.add(str(resolved))
        keys.add(os.path.normcase(str(resolved)))
    except OSError:
        pass
    if isinstance(src_value, str) and src_value.startswith("file://"):
        keys.add(src_value[7:])
    return {key for key in keys if key}


def _remember_path_rewrite(
    rewrites: dict[str, str],
    src_value: Any,
    copied_value: str | None,
) -> None:
    if not copied_value:
        return
    src = _path_from_value(src_value)
    if src is None:
        return
    for key in _path_rewrite_keys(src_value, src):
        rewrites.setdefault(key, copied_value)


def _lookup_rewritten_path(value: str, rewrites: dict[str, str]) -> str | None:
    if value in rewrites:
        return rewrites[value]
    src = _path_from_value(value)
    if src is None:
        return None
    for key in _path_rewrite_keys(value, src):
        if key in rewrites:
            return rewrites[key]
    return None


def _rewrite_artifact_paths(value: Any, rewrites: dict[str, str]) -> Any:
    if not rewrites:
        return _to_jsonable(value)
    if not isinstance(value, (dict, list, tuple, set, Path, str, int, float, bool)) and value is not None:
        return _rewrite_artifact_paths(_to_jsonable(value), rewrites)
    if isinstance(value, dict):
        return {str(key): _rewrite_artifact_paths(item, rewrites) for key, item in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_rewrite_artifact_paths(item, rewrites) for item in value]
    if isinstance(value, Path):
        return _lookup_rewritten_path(str(value), rewrites) or str(value)
    if isinstance(value, str):
        return _lookup_rewritten_path(value, rewrites) or value
    return _to_jsonable(value)


def _copy_payload_referenced_artifacts(
    *,
    step_dir: Path,
    payload: Any,
    rewrites: dict[str, str],
) -> None:
    image_suffixes = {".png", ".jpg", ".jpeg"}
    xml_suffixes = {".xml"}
    seen: set[str] = set()

    def _walk(value: Any) -> None:
        if isinstance(value, dict):
            for item in value.values():
                _walk(item)
            return
        if isinstance(value, (list, tuple, set)):
            for item in value:
                _walk(item)
            return
        if not isinstance(value, (str, os.PathLike)):
            return
        src = _path_from_value(value)
        if src is None:
            return
        try:
            resolved = src.resolve()
        except OSError:
            resolved = src
        key = os.path.normcase(str(resolved))
        if key in seen:
            return
        seen.add(key)
        if any(rewrite_key in rewrites for rewrite_key in _path_rewrite_keys(value, src)):
            return
        try:
            resolved.relative_to(step_dir.resolve())
            _remember_path_rewrite(rewrites, value, str(src))
            return
        except ValueError:
            pass
        suffix = src.suffix.lower()
        if suffix in image_suffixes:
            copied = _copy_step_artifact(value, step_dir / "screenshots", prefix="payload_")
        elif suffix in xml_suffixes:
            copied = _copy_step_artifact(value, step_dir / "xml", prefix="payload_")
        else:
            copied = None
        _remember_path_rewrite(rewrites, value, copied)

    _walk(_to_jsonable(payload))


def _messages_to_text(messages: Any) -> str:
    if not messages:
        return ""
    if isinstance(messages, list):
        chunks: list[str] = []
        for message in messages:
            if isinstance(message, dict):
                role = message.get("role") or "message"
                content = message.get("content")
                chunks.append(f"[{role}]\n{_truncate_text(json.dumps(_to_jsonable(content), ensure_ascii=False) if not isinstance(content, str) else content)}")
            else:
                chunks.append(_truncate_text(str(message)))
        return "\n\n".join(chunks)
    if isinstance(messages, str):
        return messages
    return json.dumps(_to_jsonable(messages), ensure_ascii=False, indent=2)


def _parse_vlm_call_repr(value: Any) -> dict[str, Any] | None:
    if isinstance(value, dict):
        return value
    if not isinstance(value, str):
        return None
    text = value.strip()
    if not text.startswith("VLMCallingData(") or not text.endswith(")"):
        return None
    try:
        parsed = ast.parse(text, mode="eval")
    except SyntaxError:
        return None
    call = parsed.body
    if not isinstance(call, ast.Call):
        return None
    func = call.func
    if not isinstance(func, ast.Name) or func.id != "VLMCallingData":
        return None
    output: dict[str, Any] = {}
    for keyword in call.keywords:
        if keyword.arg is None:
            continue
        try:
            output[keyword.arg] = ast.literal_eval(keyword.value)
        except Exception:
            output[keyword.arg] = ast.unparse(keyword.value) if hasattr(ast, "unparse") else None
    return output


def _model_identity_from_call(call: dict[str, Any]) -> dict[str, Any]:
    """Preserve requested and provider-resolved model ids in report ledgers."""
    requested = call.get("requested_model_name")
    resolved = call.get("resolved_model_name")
    legacy = call.get("model_name")
    return {
        "model_name": resolved or legacy or requested,
        "requested_model_name": requested,
        "resolved_model_name": resolved,
        "resolved_model_source": call.get("resolved_model_source"),
        "resolved_model_missing_reason": call.get("resolved_model_missing_reason"),
    }


def _extract_mobileuse_episode_fields(value: Any) -> dict[str, Any]:
    if isinstance(value, dict):
        return value
    if not isinstance(value, str):
        return {}
    text = value.strip()
    if not text.startswith("MobileUseEpisodeData(") or not text.endswith(")"):
        return {}
    output: dict[str, Any] = {}
    goal_match = re.search(r"goal=(?P<quote>['\"])(?P<goal>(?:\\.|(?!\1).)*)\1,\s*num_steps=", text, re.DOTALL)
    if goal_match:
        raw_goal = goal_match.group("goal")
        try:
            output["goal"] = ast.literal_eval(f"'{raw_goal}'")
        except Exception:
            output["goal"] = raw_goal.replace("\\'", "'").replace('\\"', '"')
    return output


def _llm_rows_by_step(llm_rows: Sequence[dict[str, Any]]) -> dict[int, list[dict[str, Any]]]:
    grouped: dict[int, list[dict[str, Any]]] = {}
    for row in llm_rows:
        try:
            step = int(row.get("step") or 0)
        except (TypeError, ValueError):
            continue
        if step <= 0:
            continue
        grouped.setdefault(step, []).append(row)
    return grouped


def _trace_rows_from_llm_rows(llm_rows: Sequence[dict[str, Any]]) -> list[dict[str, Any]]:
    trace_rows: list[dict[str, Any]] = []
    for step_index, rows in sorted(_llm_rows_by_step(llm_rows).items()):
        if not rows:
            continue
        last = rows[-1]
        metadata = {
            "summary": last.get("summary") or last.get("action") or last.get("response"),
            "action": last.get("action"),
        }
        current_response = {
            "llm_calls": rows,
            "metadata": metadata,
            "step_duration": last.get("step_duration"),
            "usage": last.get("usage"),
            "raw_response": last.get("response"),
        }
        trace_rows.append(
            {
                "step": step_index,
                "current_response": current_response,
                "parsed_action": last.get("action"),
                "image": last.get("screenshot"),
                "xml": last.get("xml"),
                "recovered_from_llm_log": True,
            }
        )
    return trace_rows


def _write_step_llm_artifacts(step_dir: Path, llm_rows: Sequence[dict[str, Any]]) -> dict[str, str]:
    calls_dir = step_dir / "llm_calls"
    screenshots_dir = step_dir / "screenshots"
    xml_dir = step_dir / "xml"
    path_rewrites: dict[str, str] = {}
    if not llm_rows:
        _write_jsonl(calls_dir / "llm_calls.jsonl", [])
        _write_json(
            calls_dir / "no_llm_call.json",
            {
                "llm_call_count": 0,
                "usage": None,
                "missing_reason": "no_llm_calls_recorded",
            },
        )
        _write_text(
            calls_dir / "no_llm_call_prompt.txt",
            "本步骤没有记录到模型提示词；原始事实：该步骤没有持久化 LLM call。\n",
        )
        _write_text(
            calls_dir / "no_llm_call_response.txt",
            "本步骤没有记录到模型返回；原始事实：该步骤没有持久化 LLM response。\n",
        )
        return path_rewrites
    rewritten_rows: list[dict[str, Any]] = []
    for call_index, row in enumerate(llm_rows, start=1):
        call_prefix = f"call_{call_index:03d}"
        copied = _copy_step_artifact(row.get("screenshot"), screenshots_dir, prefix=f"{call_prefix}_")
        _remember_path_rewrite(path_rewrites, row.get("screenshot"), copied)
        for image_index, image in enumerate(row.get("screenshots") or [], start=1):
            copied = _copy_step_artifact(image, screenshots_dir, prefix=f"{call_prefix}_image_{image_index}_")
            _remember_path_rewrite(path_rewrites, image, copied)
        copied = _copy_step_artifact(row.get("xml"), xml_dir, prefix=f"{call_prefix}_")
        _remember_path_rewrite(path_rewrites, row.get("xml"), copied)
        for xml_index, xml in enumerate(row.get("xmls") or [], start=1):
            copied = _copy_step_artifact(xml, xml_dir, prefix=f"{call_prefix}_xml_{xml_index}_")
            _remember_path_rewrite(path_rewrites, xml, copied)

        rewritten_row = _rewrite_artifact_paths(row, path_rewrites)
        if not isinstance(rewritten_row.get("usage_normalized"), dict):
            rewritten_row["usage_normalized"] = normalize_usage(rewritten_row.get("usage"))
        rewritten_rows.append(rewritten_row)
        _write_json(calls_dir / f"{call_prefix}.json", rewritten_row)
        prompt = row.get("prompt")
        messages = rewritten_row.get("messages")
        if prompt is None and messages is not None:
            prompt = _messages_to_text(messages)
        if prompt is not None:
            _write_text(calls_dir / f"{call_prefix}_prompt.txt", str(prompt))
        if messages is not None:
            _write_json(calls_dir / f"{call_prefix}_messages.json", messages)
        response = rewritten_row.get("response")
        if response is not None:
            _write_text(calls_dir / f"{call_prefix}_response.txt", str(response))
    _write_jsonl(calls_dir / "llm_calls.jsonl", rewritten_rows)
    return path_rewrites


def _extract_knowledge(step_payload: Any, llm_rows: Sequence[dict[str, Any]]) -> list[Any]:
    payload = _to_jsonable(step_payload)
    candidates: list[Any] = []
    if isinstance(payload, dict):
        for key in ("knowledge", "retrieved_knowledge", "explored_knowledge"):
            if payload.get(key) is not None:
                candidates.append(payload.get(key))
        metadata = payload.get("metadata")
        if isinstance(metadata, dict):
            for key in ("knowledge", "retrieved_knowledge", "explored_knowledge"):
                if metadata.get(key) is not None:
                    candidates.append(metadata.get(key))
        response = payload.get("current_response")
        if isinstance(response, dict):
            for key in ("knowledge", "retrieved_knowledge", "explored_knowledge"):
                if response.get(key) is not None:
                    candidates.append(response.get(key))
    for row in llm_rows:
        if row.get("knowledge") is not None:
            candidates.append(row.get("knowledge"))
    flattened: list[Any] = []
    for item in candidates:
        if isinstance(item, list):
            flattened.extend(item)
        else:
            flattened.append(item)
    deduped: list[Any] = []
    seen: set[str] = set()
    for item in flattened:
        key = json.dumps(_to_jsonable(item), ensure_ascii=False, sort_keys=True)
        if key in seen:
            continue
        seen.add(key)
        deduped.append(item)
    return deduped


def _extract_error(step_payload: Any) -> str | None:
    payload = _to_jsonable(step_payload)
    if not isinstance(payload, dict):
        return None
    for key in ("error", "reflection_error", "trajectory_reflection_error", "evaluation_reason"):
        value = payload.get(key)
        if value:
            return str(value)
    response = payload.get("current_response")
    if isinstance(response, dict):
        for key in ("error", "reflection_error", "trajectory_reflection_error", "evaluation_reason"):
            value = response.get(key)
            if value:
                return str(value)
    return None


def _write_step_common_artifacts(
    *,
    report_dir: Path,
    step_index: int,
    step_payload: Any,
    llm_rows: Sequence[dict[str, Any]],
) -> Path:
    step_dir = report_dir / "steps" / f"step_{step_index:03d}"
    step_dir.mkdir(parents=True, exist_ok=True)
    path_rewrites = _write_step_llm_artifacts(step_dir, llm_rows)
    _copy_payload_referenced_artifacts(
        step_dir=step_dir,
        payload=step_payload,
        rewrites=path_rewrites,
    )
    rewritten_step_payload = _rewrite_artifact_paths(step_payload, path_rewrites)
    _write_json(step_dir / "step.json", rewritten_step_payload)
    _write_json(step_dir / "knowledge.json", {"knowledge": _extract_knowledge(rewritten_step_payload, llm_rows)})
    error = _extract_error(rewritten_step_payload)
    if error:
        _write_text(step_dir / "error.txt", error)
    return step_dir


def _env_state_dom_tree(env_state: Any) -> str | None:
    if env_state is None:
        return None
    if isinstance(env_state, dict):
        dom_tree = env_state.get("dom_tree")
    else:
        dom_tree = getattr(env_state, "dom_tree", None)
    if dom_tree is None:
        return None
    if isinstance(dom_tree, bytes):
        dom_tree = dom_tree.decode("utf-8", errors="replace")
    text = str(dom_tree).strip()
    return text or None


def _write_env_state_dom_tree(env_state: Any, xml_dir: Path, filename: str) -> None:
    dom_tree = _env_state_dom_tree(env_state)
    if dom_tree is None:
        return
    _write_text(xml_dir / filename, dom_tree)


def _explicit_step_captures(step: Any) -> dict[str, Any] | None:
    """Return captures emitted by a runtime, without inferring neighboring state."""
    payload = _to_jsonable(step)
    if not isinstance(payload, dict):
        return None
    explicit = payload.get("state_captures")
    if isinstance(explicit, dict) and any(key in explicit for key in ("before", "after")):
        return explicit
    boundary = payload.get("action_boundary_captures")
    if not isinstance(boundary, list):
        return None
    before = next(
        (item for item in boundary if isinstance(item, dict) and item.get("phase") == "before"),
        None,
    )
    after = next(
        (item for item in reversed(boundary) if isinstance(item, dict) and item.get("phase") == "after"),
        None,
    )
    if before is None and after is None:
        return None
    return {
        "before": before,
        "after": after,
        "capture_source": "action_boundary_captures",
    }


def _write_androidworld_step_dirs(
    report_dir: Path,
    step_list: Sequence[Any],
    llm_rows: Sequence[dict[str, Any]],
    *,
    strict_capture: bool = False,
) -> None:
    grouped = _llm_rows_by_step(llm_rows)
    for index, step in enumerate(step_list, start=1):
        step_llm_rows = grouped.get(index, [])
        if not step_llm_rows:
            step_payload = _to_jsonable(step)
            if isinstance(step_payload, dict):
                direct_llm_rows = step_payload.get("llm_calls") or step_payload.get("vlm_call_history") or []
                if isinstance(direct_llm_rows, list):
                    step_llm_rows = [
                        parsed
                        for row in direct_llm_rows
                        if (parsed := _parse_vlm_call_repr(row)) is not None
                    ]
        step_dir = _write_step_common_artifacts(
            report_dir=report_dir,
            step_index=index,
            step_payload=step,
            llm_rows=step_llm_rows,
        )
        screenshots_dir = step_dir / "screenshots"
        xml_dir = step_dir / "xml"
        if isinstance(step, dict):
            current_state = step.get("curr_env_state")
            exec_state = step.get("exec_env_state")
        else:
            current_state = getattr(step, "curr_env_state", None)
            exec_state = getattr(step, "exec_env_state", None)
        if strict_capture:
            # _write_step_common_artifacts has already copied payload-referenced
            # PNG/XML files into this step directory and rewritten their paths
            # in step.json.  Read that rewritten payload so canonical captures
            # never point at an ephemeral agent_logs location.
            try:
                rewritten_step_payload = json.loads((step_dir / "step.json").read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError):
                rewritten_step_payload = _to_jsonable(step)
            captures = _explicit_step_captures(rewritten_step_payload)
            if captures is None:
                captures = {
                    "before": capture_state_artifacts(
                        current_state,
                        report_dir / "runtime_captures",
                        step_index=index,
                        phase="before",
                    ),
                    "after": capture_state_artifacts(
                        exec_state,
                        report_dir / "runtime_captures",
                        step_index=index,
                        phase="after",
                    ),
                }
            else:
                captures = dict(captures)
                for phase in ("before", "after"):
                    if not isinstance(captures.get(phase), dict):
                        captures[phase] = capture_state_artifacts(
                            None,
                            report_dir / "runtime_captures",
                            step_index=index,
                            phase=phase,
                        )
            _write_json(step_dir / "state_captures.json", captures)
            step_payload = rewritten_step_payload
            if isinstance(step_payload, dict):
                step_payload["state_captures"] = captures
                _write_json(step_dir / "step.json", step_payload)
            # New runs must never infer a state from a filename or a nearby step.
            continue

        current_pixels = getattr(current_state, "pixels", None)
        if current_pixels is not None:
            screenshots_dir.mkdir(parents=True, exist_ok=True)
            current_pixels.save(screenshots_dir / "before.png")
        _write_env_state_dom_tree(current_state, xml_dir, "before.xml")
        exec_pixels = getattr(exec_state, "pixels", None)
        if exec_pixels is not None:
            screenshots_dir.mkdir(parents=True, exist_ok=True)
            exec_pixels.save(screenshots_dir / "after.png")
        _write_env_state_dom_tree(exec_state, xml_dir, "after.xml")
        _copy_report_step_matches(report_dir, index, step_dir)
        _copy_nearest_report_screenshot(report_dir, index, step_dir)


def _copy_report_step_matches(report_dir: Path, step_index: int, step_dir: Path) -> None:
    screenshot_patterns = [
        f"*step_{step_index:02d}*",
        f"*step_{step_index}_*",
        f"step_{step_index}.*",
        f"step_{step_index}_*.*",
        f"screenshot-{step_index - 1}-*",
    ]
    xml_patterns = [
        f"*step_{step_index:02d}*",
        f"*step_{step_index}_*",
        f"step_{step_index}.*",
        f"screenshot_{step_index - 1}.xml",
        f"screenshot_{step_index}.xml",
        f"{step_index - 1}.xml",
    ]
    for pattern in screenshot_patterns:
        for path in sorted((report_dir / "screenshots").glob(pattern), key=lambda item: _natural_key(item.name)):
            _copy_step_artifact(path, step_dir / "screenshots", prefix="report_")
    for pattern in xml_patterns:
        for path in sorted((report_dir / "xml").glob(pattern), key=lambda item: _natural_key(item.name)):
            _copy_step_artifact(path, step_dir / "xml", prefix="report_")


def _step_number_from_artifact_name(path: Path) -> int | None:
    match = re.search(r"step[_-](\d+)", path.name)
    if match:
        try:
            return int(match.group(1))
        except ValueError:
            return None
    match = re.search(r"screenshot-(\d+)-", path.name)
    if match:
        try:
            return int(match.group(1)) + 1
        except ValueError:
            return None
    return None


def _copy_nearest_report_screenshot(report_dir: Path, step_index: int, step_dir: Path) -> None:
    screenshots_dir = step_dir / "screenshots"
    if list(screenshots_dir.glob("*.png")) or list(screenshots_dir.glob("*.jpg")) or list(screenshots_dir.glob("*.jpeg")):
        return
    candidates: list[tuple[int, Path]] = []
    for image in [
        *sorted((report_dir / "screenshots").glob("*.png"), key=lambda item: _natural_key(item.name)),
        *sorted((report_dir / "screenshots").glob("*.jpg"), key=lambda item: _natural_key(item.name)),
        *sorted((report_dir / "screenshots").glob("*.jpeg"), key=lambda item: _natural_key(item.name)),
    ]:
        number = _step_number_from_artifact_name(image)
        if number is None:
            continue
        candidates.append((number, image))
    previous_or_equal = [(number, image) for number, image in candidates if number <= step_index]
    if previous_or_equal:
        _, image = sorted(previous_or_equal, key=lambda item: (item[0], _natural_key(item[1].name)))[-1]
        _copy_step_artifact(image, screenshots_dir, prefix="fallback_nearest_")
        return
    fallback_images = [
        *sorted((report_dir / "screenshots").glob("*.png"), key=_safe_mtime),
        *sorted((report_dir / "screenshots").glob("*.jpg"), key=_safe_mtime),
        *sorted((report_dir / "screenshots").glob("*.jpeg"), key=_safe_mtime),
    ]
    if fallback_images:
        _copy_step_artifact(fallback_images[-1], screenshots_dir, prefix="fallback_latest_")


def _write_generic_step_dirs(
    report_dir: Path,
    trace_rows: Sequence[dict[str, Any]],
    llm_rows: Sequence[dict[str, Any]],
    *,
    strict_capture: bool = False,
) -> None:
    grouped = _llm_rows_by_step(llm_rows)
    for index, trace_row in enumerate(trace_rows, start=1):
        step_dir = _write_step_common_artifacts(
            report_dir=report_dir,
            step_index=index,
            step_payload=trace_row,
            llm_rows=grouped.get(index, []),
        )
        if not strict_capture:
            _copy_report_step_matches(report_dir, index, step_dir)
            _copy_nearest_report_screenshot(report_dir, index, step_dir)


def _write_androidlab_step_dirs(
    report_dir: Path,
    trace_rows: Sequence[dict[str, Any]],
    llm_rows: Sequence[dict[str, Any]],
) -> None:
    grouped = _llm_rows_by_step(llm_rows)
    for index, trace_row in enumerate(trace_rows, start=1):
        step_dir = _write_step_common_artifacts(
            report_dir=report_dir,
            step_index=index,
            step_payload=trace_row,
            llm_rows=grouped.get(index, []),
        )
        screenshots_dir = step_dir / "screenshots"
        xml_dir = step_dir / "xml"
        _copy_step_artifact(trace_row.get("image"), screenshots_dir, prefix="record_")
        _copy_step_artifact(trace_row.get("xml"), xml_dir, prefix="record_")
        _copy_step_artifact(trace_row.get("ac_xml"), xml_dir, prefix="a11y_")
        for suffix in ("curr_env", "exec_env"):
            _copy_step_artifact(
                report_dir / "screenshots" / f"step_{index}_{suffix}.png",
                screenshots_dir,
                prefix="state_",
            )
        _copy_step_artifact(report_dir / "xml" / f"{index - 1}.xml", xml_dir, prefix="state_")
        _copy_report_step_matches(report_dir, index, step_dir)


def _backfill_step_dirs_from_task_reports(run_output_dir: Path, benchmark: str) -> None:
    for report_dir in sorted(run_output_dir.glob("task_reports/*"), key=lambda item: _natural_key(item.name)):
        if (report_dir / "steps").exists():
            continue
        trace_rows = list(_iter_jsonl(report_dir / "logs" / "step_trace.jsonl"))
        llm_rows = list(_iter_jsonl(report_dir / "logs" / "llm_interactions.jsonl"))
        if not trace_rows and llm_rows:
            trace_rows = _trace_rows_from_llm_rows(llm_rows)
            _write_jsonl(report_dir / "logs" / "step_trace.jsonl", trace_rows)
        if not trace_rows:
            continue
        if benchmark == "androidlab":
            _write_androidlab_step_dirs(report_dir, trace_rows, llm_rows)
        else:
            _write_generic_step_dirs(report_dir, trace_rows, llm_rows)


def _write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def _should_localize_large_run_files() -> bool:
    return os.getenv("STUDY_LOCALIZE_RUN_FILES", "").strip().lower() in {"1", "true", "yes", "on"}


def _truncate_text(text: str, limit: int = 30000) -> str:
    if len(text) <= limit:
        return text
    return text[: limit - 3] + "..."


def _find_outputs_root(path: Path) -> Path:
    for candidate in [path, *path.parents]:
        if candidate.name == "outputs":
            return candidate
    raise ValueError(f"Could not find outputs root for {path}")


def _parse_screenshot_epoch(path: Path) -> float | None:
    match = re.search(r"screenshot-\d+-([0-9]+(?:\.[0-9]+)?)-", path.name)
    if not match:
        return None
    return float(match.group(1))


def _format_seconds(value: float | None) -> float | None:
    if value is None:
        return None
    return round(float(value), 3)


def _coerce_success(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return False
        return bool(value)
    if isinstance(value, str):
        return value.strip().lower() in {"1", "true", "yes", "success", "successful"}
    return False


def _ensure_workbook(workbook_path: Path) -> tuple[Workbook, Any]:
    if Workbook is None or load_workbook is None:
        raise RuntimeError(
            "Legacy benchmark Excel export requires openpyxl; use "
            "study/scripts/export_runtime_metrics_excel.mjs for runtime workbooks."
        )
    if workbook_path.exists():
        workbook = load_workbook(workbook_path)
        sheet = workbook.active
        if sheet.max_row == 0:
            sheet.append([EXCEL_HEADERS_ZH[column] for column in EXCEL_COLUMNS])
        elif sheet.max_row == 1 and sheet.cell(1, 1).value is None:
            sheet.delete_rows(1, 1)
            sheet.append([EXCEL_HEADERS_ZH[column] for column in EXCEL_COLUMNS])
        return workbook, sheet

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "study_results"
    sheet.append([EXCEL_HEADERS_ZH[column] for column in EXCEL_COLUMNS])
    return workbook, sheet


def write_records_to_excel(workbook_path: Path, records: Sequence[dict[str, Any]]) -> Path:
    if Workbook is None:
        raise RuntimeError(
            "Legacy benchmark Excel export requires openpyxl; use "
            "study/scripts/export_runtime_metrics_excel.mjs for runtime workbooks."
        )
    translator = ChineseTranslator()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "study_results"
    sheet.append([EXCEL_HEADERS_ZH[column] for column in EXCEL_COLUMNS])
    for record in records:
        localized = localize_record(record, translator)
        sheet.append([localized.get(column) for column in EXCEL_COLUMNS])
    workbook.save(workbook_path)
    return workbook_path


def append_records_to_excel(workbook_path: Path, records: Sequence[dict[str, Any]]) -> Path:
    if Workbook is None:
        raise RuntimeError(
            "Legacy benchmark Excel export requires openpyxl; use "
            "study/scripts/export_runtime_metrics_excel.mjs for runtime workbooks."
        )
    translator = ChineseTranslator()
    workbook, sheet = _ensure_workbook(workbook_path)
    row_index_by_record_id: dict[str, int] = {}
    for row_idx in range(2, sheet.max_row + 1):
        record_id = sheet.cell(row_idx, 1).value
        if record_id:
            row_index_by_record_id[str(record_id)] = row_idx

    for record in records:
        record_id = str(record["record_id"])
        localized = localize_record(record, translator)
        row_idx = row_index_by_record_id.get(record_id)
        if row_idx is None:
            row_idx = sheet.max_row + 1
            row_index_by_record_id[record_id] = row_idx
        for col_idx, column in enumerate(EXCEL_COLUMNS, start=1):
            sheet.cell(row_idx, col_idx).value = localized.get(column)

    workbook.save(workbook_path)
    return workbook_path


def append_records_to_index(index_path: Path, records: Sequence[dict[str, Any]]) -> Path:
    translator = ChineseTranslator()
    existing = _load_json(index_path)
    existing_records = list(existing) if isinstance(existing, list) else []
    by_record_id = {str(record.get("record_id")): record for record in existing_records if record.get("record_id")}
    for record in records:
        by_record_id[str(record["record_id"])] = localize_record(record, translator)
    merged = sorted(by_record_id.values(), key=lambda item: _natural_key(str(item["record_id"])))
    _write_json(index_path, merged)
    return index_path


def benchmark_report_paths(outputs_root: str | Path, benchmark: str) -> dict[str, Path]:
    outputs_root = Path(outputs_root).resolve()
    normalized = str(benchmark).lower().replace("-", "")
    if normalized not in {"androidworld", "androidlab"}:
        normalized = str(benchmark).lower()
    return {
        "excel": outputs_root / f"{normalized}_results.xlsx",
        "index": outputs_root / f"{normalized}_report_index.json",
    }


def append_records_to_benchmark_outputs(outputs_root: str | Path, benchmark: str, records: Sequence[dict[str, Any]]) -> dict[str, Path]:
    paths = benchmark_report_paths(outputs_root, benchmark)
    append_records_to_excel(paths["excel"], records)
    append_records_to_index(paths["index"], records)
    return paths


def _build_androidlab_task_catalog() -> dict[str, dict[str, Any]]:
    catalog: dict[str, dict[str, Any]] = {}
    config_dir = PROJECT_ROOT / "Android-Lab" / "evaluation" / "config"
    for config_path in sorted(config_dir.glob("*.yaml")):
        app_config = AppConfig(str(config_path))
        for task_id, task_name in app_config.task_name.items():
            catalog[task_id] = {
                "task": task_name.strip(),
                "task_config_path": str(config_path),
                "app": getattr(app_config, "APP", ""),
                "package": getattr(app_config, "package", ""),
                "metrics": app_config.metrics,
            }
    return catalog


_ANDROIDLAB_TASK_CATALOG: dict[str, dict[str, Any]] | None = None


def _get_androidlab_task_catalog() -> dict[str, dict[str, Any]]:
    global _ANDROIDLAB_TASK_CATALOG
    if _ANDROIDLAB_TASK_CATALOG is None:
        _ANDROIDLAB_TASK_CATALOG = _build_androidlab_task_catalog()
    return _ANDROIDLAB_TASK_CATALOG


def _evaluate_androidlab_task(task_id: str, task_dir: Path, task_catalog: dict[str, dict[str, Any]]) -> bool:
    task_meta = task_catalog.get(task_id)
    if task_meta is None:
        return False

    app_config = AppConfig(task_meta["task_config_path"])
    metric_factory = app_config.metrics.get(task_id)
    if metric_factory is None:
        return False

    trace_file = task_dir / "traces" / "trace.jsonl"
    xml_dir = task_dir / "xml"
    if not trace_file.exists() or not xml_dir.exists():
        return False

    metric = metric_factory(type("Args", (), {"judge_model": "glm4", "api_base": "", "api_key": ""})())
    final_result = {"complete": False}
    last_action = None
    repeat_count = 0

    with trace_file.open("r", encoding="utf-8") as handle:
        for raw_line in handle:
            line = json.loads(raw_line)
            current_action = json.dumps(line.get("parsed_action", {}), ensure_ascii=False, sort_keys=True)
            if current_action == last_action:
                repeat_count += 1
                if repeat_count > 5:
                    break
            else:
                repeat_count = 0
                last_action = current_action

            xml_name = line.get("ac_xml") or line.get("xml")
            if not xml_name:
                continue
            xml_path = xml_dir / Path(xml_name).name
            if not xml_path.exists():
                continue

            xml_compressed = dump_xml(str(xml_path))
            if xml_compressed is None:
                continue

            try:
                result = metric.judge(xml_compressed, line)
            except Exception:
                continue
            if "judge_page" in result and not result.get("judge_page"):
                continue
            final_result = result

    return bool(final_result.get("complete"))


def _iter_jsonl(path: Path) -> Iterable[dict[str, Any]]:
    if not path.exists():
        return []
    rows: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as handle:
        for raw_line in handle:
            raw_line = raw_line.strip()
            if not raw_line:
                continue
            rows.append(json.loads(raw_line))
    return rows


def _usage_from_payload(payload: Any) -> dict[str, Any]:
    if isinstance(payload, dict) and "usage_present" in payload:
        return dict(payload)
    return normalize_usage(payload)


def _sum_usage_from_rows(rows: Iterable[dict[str, Any]]) -> dict[str, int]:
    calls: list[dict[str, Any]] = []
    for row in rows:
        if not isinstance(row, dict):
            continue

        def _add_usage(usage_payload: Any) -> None:
            calls.append({"usage": usage_payload})

        direct_calls = row.get("llm_calls") or row.get("vlm_call_history") or []
        if direct_calls:
            for call in direct_calls:
                parsed_call = _parse_vlm_call_repr(call)
                if parsed_call is not None:
                    _add_usage(parsed_call.get("usage"))
            continue

        current_response = row.get("current_response")
        if isinstance(current_response, dict):
            nested_calls = current_response.get("llm_calls") or current_response.get("vlm_call_history") or []
            if nested_calls:
                for call in nested_calls:
                    if isinstance(call, dict):
                        _add_usage(call.get("usage"))
                continue
            for key in ("usage", "memory_usage", "reflect_usage", "planning_usage"):
                if current_response.get(key) is not None:
                    _add_usage(current_response.get(key))
            if any(current_response.get(key) is not None for key in ("usage", "memory_usage", "reflect_usage", "planning_usage")):
                continue

        for key in ("usage", "memory_usage", "reflect_usage", "planning_usage"):
            if row.get(key) is not None:
                _add_usage(row.get(key))
    return aggregate_usage(calls)


def _sum_usage_by_stage_from_rows(rows: Iterable[dict[str, Any]]) -> dict[str, dict[str, int]]:
    grouped: dict[str, list[dict[str, Any]]] = {}

    def _add(stage: Any, usage_payload: Any) -> None:
        stage_name = str(stage or "unknown")
        grouped.setdefault(stage_name, []).append({"usage": usage_payload})

    for row in rows:
        if not isinstance(row, dict):
            continue
        direct_calls = row.get("llm_calls") or row.get("vlm_call_history") or []
        if direct_calls:
            for call in direct_calls:
                parsed_call = _parse_vlm_call_repr(call)
                if parsed_call is not None:
                    _add(parsed_call.get("stage") or row.get("stage"), parsed_call.get("usage"))
            continue

        current_response = row.get("current_response")
        if isinstance(current_response, dict):
            nested_calls = current_response.get("llm_calls") or current_response.get("vlm_call_history") or []
            if nested_calls:
                for call in nested_calls:
                    parsed_call = _parse_vlm_call_repr(call)
                    if parsed_call is not None:
                        _add(parsed_call.get("stage") or current_response.get("stage"), parsed_call.get("usage"))
                continue
            for key in ("usage", "memory_usage", "reflect_usage", "planning_usage"):
                if current_response.get(key) is not None:
                    _add(current_response.get("stage") or key.replace("_usage", ""), current_response.get(key))
            continue

        for key in ("usage", "memory_usage", "reflect_usage", "planning_usage"):
            if row.get(key) is not None:
                _add(row.get("stage") or key.replace("_usage", ""), row.get(key))
    return {stage: aggregate_usage(calls) for stage, calls in grouped.items()}


def _collect_mobileuse_usage_from_steps(step_list: Sequence[Any]) -> dict[str, int]:
    calls: list[Any] = []
    for step in step_list:
        for call in getattr(step, "vlm_call_history", []) or []:
            calls.append(call)
    return aggregate_usage(calls)


def _usage_has_known_tokens(usage: dict[str, Any] | None) -> bool:
    return bool(usage) and any(isinstance(usage.get(key), int) for key in ("prompt_tokens", "completion_tokens", "total_tokens"))


def _extract_step_duration(step: Any) -> float | None:
    for attr in ("step_duration", "exec_duration"):
        value = getattr(step, attr, None)
        if isinstance(value, (int, float)):
            return float(value)
    if isinstance(step, dict):
        for attr in ("step_duration", "exec_duration"):
            value = step.get(attr)
            if isinstance(value, (int, float)):
                return float(value)
    return None


def _discover_agent_dirs(*roots: Path) -> list[Path]:
    candidates: list[Path] = []
    for root in roots:
        if not root.exists():
            continue
        for child in sorted(root.iterdir(), key=lambda item: _natural_key(item.name)):
            if not child.is_dir() or child.name.startswith("_"):
                continue
            # Mobile-Agent-v3 stores all tasks below one shared trajectory
            # container.  Returning that container as a single candidate makes
            # report generation copy the complete run into every task report.
            # Expose the task directories instead, while preserving the prior
            # behavior for other agent layouts.
            if child.name == "mobileagentv3_traj":
                task_dirs = [
                    task_dir
                    for task_dir in sorted(child.iterdir(), key=lambda item: _natural_key(item.name))
                    if task_dir.is_dir() and not task_dir.name.startswith("_")
                ]
                if task_dirs:
                    candidates.extend(task_dirs)
                    continue
            candidates.append(child)
    return candidates


def _match_agent_dir(candidates: Sequence[Path], *, task: str, task_id: str) -> Path | None:
    if not candidates:
        return None
    if len(candidates) == 1:
        return candidates[0]

    task_slug = sanitize_name(task, limit=40).lower()
    task_id_lower = task_id.lower()
    best_candidate = None
    best_score = -1
    for candidate in candidates:
        score = 0
        lowered = candidate.name.lower()
        if task_slug and task_slug in lowered:
            score += 100
        if task_id_lower and task_id_lower in lowered:
            score += 50
        episode_path = candidate / "episode.json"
        if episode_path.exists():
            episode = _extract_mobileuse_episode_fields(_load_json(episode_path)) or {}
            if str(episode.get("goal", "")).strip() == task.strip():
                score += 200
        if score > best_score:
            best_score = score
            best_candidate = candidate
    return best_candidate


def _iter_nested_strings(value: Any) -> Iterable[str]:
    if isinstance(value, str):
        yield value
    elif isinstance(value, dict):
        for item in value.values():
            yield from _iter_nested_strings(item)
    elif isinstance(value, (list, tuple)):
        for item in value:
            yield from _iter_nested_strings(item)


def _agent_dir_from_step_artifacts(
    candidates: Sequence[Path],
    step_list: Sequence[Any],
) -> Path | None:
    """Bind a checkpoint to the AppAgent directory referenced by its captures.

    P2 task text is not unique, so task-name matching cannot establish data
    provenance.  Screenshot/XML paths embedded in the checkpoint are emitted
    by the exact AppAgent session and provide an unambiguous directory link.
    """
    if not candidates or not step_list:
        return None
    resolved_candidates = [(candidate, candidate.resolve()) for candidate in candidates]
    reference_counts: dict[Path, int] = {}
    for step in step_list:
        payload = _to_jsonable(step)
        for raw_value in _iter_nested_strings(payload):
            if not raw_value or raw_value.startswith("<"):
                continue
            try:
                artifact_path = Path(raw_value).expanduser().resolve(strict=False)
            except (OSError, RuntimeError, ValueError):
                continue
            for candidate, resolved_candidate in resolved_candidates:
                try:
                    artifact_path.relative_to(resolved_candidate)
                except ValueError:
                    continue
                reference_counts[candidate] = reference_counts.get(candidate, 0) + 1
    if not reference_counts:
        return None
    return max(
        reference_counts,
        key=lambda candidate: (reference_counts[candidate], -candidates.index(candidate)),
    )


def _select_runtime_llm_rows(
    step_list: Sequence[Any],
    runtime_log_rows: Sequence[dict[str, Any]],
) -> tuple[list[dict[str, Any]], str]:
    """Choose the checkpoint ledger first because it is task/step scoped."""
    checkpoint_rows = _build_runtime_llm_rows_from_step_list(step_list)
    if checkpoint_rows:
        return checkpoint_rows, "checkpoint_step_data"
    if runtime_log_rows:
        return list(runtime_log_rows), "matched_agent_log_fallback"
    return [], "missing"


def _terminal_log_path(run_output_dir: Path) -> Path | None:
    terminal_log = run_output_dir / "logs" / "terminal.log"
    if terminal_log.exists():
        return terminal_log
    return None


def _step_summary_from_mobileuse(step: Any, index: int) -> str:
    action_desc = getattr(step, "action_desc", None)
    if action_desc:
        return f"{index}. {action_desc}"
    summary = getattr(step, "summary", None)
    if summary:
        return f"{index}. {summary}"
    action = getattr(step, "action", None)
    if action is not None:
        return f"{index}. {action}"
    return f"{index}. step"


def _step_summary_from_runtime_step(step: Any, index: int) -> str:
    if isinstance(step, dict):
        for key in ("action_desc", "summary", "action_s", "content"):
            value = step.get(key)
            if value:
                return f"{index}. {_truncate_text(str(value), 800)}"
        metadata = step.get("metadata")
        if isinstance(metadata, dict):
            summary = metadata.get("summary") or metadata.get("last_act")
            if summary:
                return f"{index}. {_truncate_text(str(summary), 800)}"
        action = step.get("action")
        if action:
            return f"{index}. {_truncate_text(json.dumps(_to_jsonable(action), ensure_ascii=False), 800)}"
        error = step.get("error")
        if error:
            return f"{index}. error: {error}"
        return f"{index}. step"
    metadata = getattr(step, "metadata", None) or {}
    summary = metadata.get("summary") or metadata.get("last_act")
    if summary:
        return f"{index}. {summary}"
    action_result = getattr(step, "action_result", None) or {}
    action = action_result.get("action")
    if action:
        return f"{index}. {action}"
    error = getattr(step, "error", None)
    if error:
        return f"{index}. error: {error}"
    return f"{index}. step"


def _step_summary_from_trace(line: dict[str, Any], index: int) -> str:
    current_response = line.get("current_response") or {}
    if isinstance(current_response, dict):
        action_desc = current_response.get("action_desc")
        if action_desc:
            return f"{index}. {action_desc}"
        metadata = current_response.get("metadata") or {}
        summary = metadata.get("summary") or metadata.get("last_act")
        if summary:
            return f"{index}. {summary}"
    parsed_action = line.get("parsed_action") or {}
    action = parsed_action.get("action")
    if action:
        return f"{index}. {action}"
    return f"{index}. step"


def _copy_terminal_log(run_output_dir: Path, report_dir: Path, notes: list[str]) -> None:
    terminal_log = _terminal_log_path(run_output_dir)
    logs_dir = report_dir / "logs"
    if terminal_log is not None:
        _copy_file(terminal_log, logs_dir / "terminal.log")
        return
    notes.append("terminal_log_missing")
    _write_text(logs_dir / "terminal.log", "Terminal log was not captured for this run.\n")


def _make_base_record(
    *,
    manifest: dict[str, Any],
    benchmark: str,
    mode: str,
    run_name: str,
    run_output_dir: Path,
    task_id: str,
    task: str,
    execution_steps: str,
    success: bool,
    final_result: str,
    step_count: int,
    usage_totals: dict[str, Any],
    usage_by_stage: dict[str, dict[str, Any]] | None,
    runtime_seconds: float | None,
    started_at: str | None,
    finished_at: str | None,
    report_dir: Path,
    notes: list[str],
    token_reconciliation: dict[str, Any] | None = None,
    runtime_reconciliation: dict[str, Any] | None = None,
    terminal_event: str | None = None,
    terminal_event_missing_reason: str | None = None,
    started_at_source: str | None = None,
    started_at_missing_reason: str | None = None,
    finished_at_source: str | None = None,
    finished_at_missing_reason: str | None = None,
) -> dict[str, Any]:
    avg_step = None
    if runtime_seconds is not None and step_count:
        avg_step = runtime_seconds / step_count
    return {
        "record_id": f"{benchmark}|{mode}|{run_name}|{task_id}",
        "benchmark": benchmark,
        "mode": mode,
        "mode_label": manifest.get("mode_label", mode),
        "run_name": run_name,
        "task_id": task_id,
        "task": task,
        "execution_steps": _truncate_text(execution_steps),
        "final_result": final_result,
        "success": bool(success),
        "step_count": step_count,
        "prompt_tokens": usage_totals.get("prompt_tokens"),
        "completion_tokens": usage_totals.get("completion_tokens"),
        "total_tokens": usage_totals.get("total_tokens"),
        "known_prompt_tokens": usage_totals.get("known_prompt_tokens", 0),
        "known_completion_tokens": usage_totals.get("known_completion_tokens", 0),
        "known_total_tokens": usage_totals.get("known_total_tokens", 0),
        "usage_status": usage_totals.get("usage_status"),
        "usage_missing_reasons": usage_totals.get("usage_missing_reasons", []),
        "token_reconciliation": token_reconciliation or {},
        "runtime_reconciliation": runtime_reconciliation or {},
        "usage_by_stage": usage_by_stage or {},
        "runtime_seconds": _format_seconds(runtime_seconds),
        "avg_step_seconds": _format_seconds(avg_step),
        "started_at": started_at,
        "started_at_source": started_at_source,
        "started_at_missing_reason": started_at_missing_reason,
        "finished_at": finished_at,
        "finished_at_source": finished_at_source,
        "finished_at_missing_reason": finished_at_missing_reason,
        "terminal_event": terminal_event,
        "terminal_event_missing_reason": terminal_event_missing_reason,
        "task_report_dir": str(report_dir),
        "run_output_dir": str(run_output_dir),
        "notes": ", ".join(notes),
    }


def _write_jsonl(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(_to_jsonable(row), ensure_ascii=False) + "\n")


def _build_mobileuse_llm_rows(step_list: Sequence[Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for index, step in enumerate(step_list, start=1):
        base_row = {
            "step": index,
            "thought": getattr(step, "thought", None),
            "action": _to_jsonable(getattr(step, "action", None)),
            "action_desc": getattr(step, "action_desc", None),
            "summary": getattr(step, "summary", None),
            "plan": getattr(step, "plan", None),
            "sub_goal": getattr(step, "sub_goal", None),
            "progress": getattr(step, "progress", None),
            "reflection_outcome": getattr(step, "reflection_outcome", None),
            "reflection_error": getattr(step, "reflection_error", None),
            "trajectory_reflection_outcome": getattr(step, "trajectory_reflection_outcome", None),
            "trajectory_reflection_error": getattr(step, "trajectory_reflection_error", None),
            "evaluation_result": getattr(step, "evaluation_result", None),
            "evaluation_reason": getattr(step, "evaluation_reason", None),
            "reflect_on_demand": getattr(step, "reflect_on_demand", None),
            "skip_reflector": getattr(step, "skip_reflector", None),
            "skip_reflector_reason": getattr(step, "skip_reflector_reason", None),
            "action_execution": _to_jsonable(getattr(step, "action_execution", None)),
            "action_executions": _to_jsonable(getattr(step, "action_executions", None)),
            "knowledge": _to_jsonable(getattr(step, "knowledge", None)),
            "step_duration": getattr(step, "step_duration", None),
            "exec_duration": getattr(step, "exec_duration", None),
        }
        call_history = _to_jsonable(getattr(step, "vlm_call_history", None) or [])
        if call_history:
            for call_index, call in enumerate(call_history, start=1):
                parsed_call = _parse_vlm_call_repr(call)
                if parsed_call is None:
                    continue
                rows.append(
                    {
                        **base_row,
                        "call_index": call_index,
                        "call_id": parsed_call.get("call_id"),
                        "stage": parsed_call.get("stage"),
                        "messages": parsed_call.get("messages"),
                        "response": parsed_call.get("response"),
                        "usage": parsed_call.get("usage"),
                        "usage_normalized": normalize_usage(parsed_call.get("usage")),
                        "role": parsed_call.get("role") or parsed_call.get("stage"),
                        **_model_identity_from_call(parsed_call),
                        "started_at": parsed_call.get("started_at"),
                        "finished_at": parsed_call.get("finished_at"),
                        "duration_seconds": parsed_call.get("duration_seconds"),
                        "error": parsed_call.get("error"),
                        "logprobs_requested": parsed_call.get("logprobs_requested"),
                        "logprobs_returned": parsed_call.get("logprobs_returned"),
                        "logprobs_missing_reason": parsed_call.get("logprobs_missing_reason"),
                    }
                )
            continue
    return rows


def _load_mobileuse_json_steps(agent_dir: Path | None) -> list[dict[str, Any]]:
    if agent_dir is None or not agent_dir.exists():
        return []
    steps: list[dict[str, Any]] = []
    for path in sorted(agent_dir.glob("step_*.json"), key=lambda item: _natural_key(item.name)):
        try:
            payload = _load_json(path)
        except Exception:
            continue
        if isinstance(payload, dict):
            steps.append(payload)
    return steps


def _build_mobileuse_llm_rows_from_json_steps(step_payloads: Sequence[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for index, payload in enumerate(step_payloads, start=1):
        base_row = {
            "step": index,
            "thought": payload.get("thought"),
            "action": _to_jsonable(payload.get("action")),
            "action_desc": payload.get("action_desc"),
            "summary": payload.get("summary"),
            "plan": payload.get("plan"),
            "sub_goal": payload.get("sub_goal"),
            "progress": payload.get("progress"),
            "reflection_outcome": payload.get("reflection_outcome"),
            "reflection_error": payload.get("reflection_error"),
            "trajectory_reflection_outcome": payload.get("trajectory_reflection_outcome"),
            "trajectory_reflection_error": payload.get("trajectory_reflection_error"),
            "evaluation_result": payload.get("evaluation_result"),
            "evaluation_reason": payload.get("evaluation_reason"),
            "reflect_on_demand": payload.get("reflect_on_demand"),
            "skip_reflector": payload.get("skip_reflector"),
            "skip_reflector_reason": payload.get("skip_reflector_reason"),
            "action_execution": _to_jsonable(payload.get("action_execution")),
            "action_executions": _to_jsonable(payload.get("action_executions")),
            "knowledge": _to_jsonable(payload.get("knowledge")),
            "step_duration": payload.get("step_duration"),
            "exec_duration": payload.get("exec_duration"),
        }
        call_history = payload.get("vlm_call_history") or []
        if call_history:
            for call_index, call in enumerate(call_history, start=1):
                parsed_call = _parse_vlm_call_repr(call)
                if parsed_call is None:
                    continue
                rows.append(
                    {
                        **base_row,
                        "call_index": call_index,
                        "call_id": parsed_call.get("call_id"),
                        "stage": parsed_call.get("stage"),
                        "messages": parsed_call.get("messages"),
                        "response": parsed_call.get("response"),
                        "usage": parsed_call.get("usage"),
                        "usage_normalized": normalize_usage(parsed_call.get("usage")),
                        "role": parsed_call.get("role") or parsed_call.get("stage"),
                        **_model_identity_from_call(parsed_call),
                        "started_at": parsed_call.get("started_at"),
                        "finished_at": parsed_call.get("finished_at"),
                        "duration_seconds": parsed_call.get("duration_seconds"),
                        "error": parsed_call.get("error"),
                        "logprobs_requested": parsed_call.get("logprobs_requested"),
                        "logprobs_returned": parsed_call.get("logprobs_returned"),
                        "logprobs_missing_reason": parsed_call.get("logprobs_missing_reason"),
                    }
                )
            continue
    return rows


def _step_summary_from_json_payload(payload: dict[str, Any], index: int) -> str:
    for key in ("action_desc", "summary", "action_s", "content"):
        value = payload.get(key)
        if value:
            return f"{index}. {_truncate_text(str(value), 800)}"
    action = payload.get("action")
    if action is not None:
        return f"{index}. {_truncate_text(json.dumps(_to_jsonable(action), ensure_ascii=False), 800)}"
    return f"{index}. step"


def _build_runtime_llm_rows_from_step_list(step_list: Sequence[Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for index, step in enumerate(step_list, start=1):
        step_payload = _to_jsonable(step)
        if not isinstance(step_payload, dict):
            continue
        llm_calls = step_payload.get("llm_calls") or step_payload.get("vlm_call_history") or []
        metadata = step_payload.get("metadata") or {}
        for call_index, call in enumerate(llm_calls, start=1):
            parsed_call = _parse_vlm_call_repr(call)
            if parsed_call is None:
                continue
            rows.append(
                {
                    "step": index,
                    "call_index": call_index,
                    "call_id": parsed_call.get("call_id"),
                    "stage": parsed_call.get("stage"),
                    "prompt": parsed_call.get("prompt"),
                    "messages": parsed_call.get("messages"),
                    "response": parsed_call.get("response"),
                    "usage": parsed_call.get("usage"),
                    "usage_normalized": normalize_usage(parsed_call.get("usage")),
                    "role": parsed_call.get("role") or parsed_call.get("stage"),
                    **_model_identity_from_call(parsed_call),
                    "started_at": parsed_call.get("started_at"),
                    "finished_at": parsed_call.get("finished_at"),
                    "duration_seconds": parsed_call.get("duration_seconds"),
                    "error": parsed_call.get("error"),
                    "logprobs_requested": parsed_call.get("logprobs_requested"),
                    "logprobs_returned": parsed_call.get("logprobs_returned"),
                    "logprobs_missing_reason": parsed_call.get("logprobs_missing_reason"),
                    "screenshot": parsed_call.get("screenshot"),
                    "screenshots": parsed_call.get("screenshots"),
                    "xml": parsed_call.get("xml"),
                    "xmls": parsed_call.get("xmls"),
                    "summary": metadata.get("summary"),
                    "action": metadata.get("action"),
                    "action_execution": step_payload.get("action_execution"),
                    "action_executions": step_payload.get("action_executions"),
                    "step_duration": step_payload.get("step_duration"),
                }
            )
    return rows


def _build_trace_llm_rows(trace_rows: Sequence[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for index, trace_row in enumerate(trace_rows, start=1):
        if not isinstance(trace_row, dict):
            continue
        current_response = trace_row.get("current_response") or {}
        if not isinstance(current_response, dict):
            continue
        llm_calls = current_response.get("llm_calls") or current_response.get("vlm_call_history") or []
        metadata = current_response.get("metadata") or {}
        for call_index, call in enumerate(llm_calls, start=1):
            parsed_call = _parse_vlm_call_repr(call)
            if parsed_call is None:
                continue
            rows.append(
                {
                    "step": index,
                    "call_index": call_index,
                    "call_id": parsed_call.get("call_id"),
                    "stage": parsed_call.get("stage"),
                    "prompt": parsed_call.get("prompt"),
                    "messages": parsed_call.get("messages"),
                    "response": parsed_call.get("response"),
                    "usage": parsed_call.get("usage"),
                    "usage_normalized": normalize_usage(parsed_call.get("usage")),
                    "role": parsed_call.get("role") or parsed_call.get("stage"),
                    **_model_identity_from_call(parsed_call),
                    "started_at": parsed_call.get("started_at"),
                    "finished_at": parsed_call.get("finished_at"),
                    "duration_seconds": parsed_call.get("duration_seconds"),
                    "error": parsed_call.get("error"),
                    "logprobs_requested": parsed_call.get("logprobs_requested"),
                    "logprobs_returned": parsed_call.get("logprobs_returned"),
                    "logprobs_missing_reason": parsed_call.get("logprobs_missing_reason"),
                    "screenshot": parsed_call.get("screenshot") or trace_row.get("image"),
                    "screenshots": parsed_call.get("screenshots"),
                    "xml": parsed_call.get("xml") or trace_row.get("ac_xml") or trace_row.get("xml"),
                    "xmls": parsed_call.get("xmls"),
                    "summary": metadata.get("summary"),
                    "action": metadata.get("action"),
                    "action_execution": current_response.get("action_execution") or trace_row.get("action_execution"),
                    "action_executions": current_response.get("action_executions") or trace_row.get("action_executions"),
                    "step_duration": current_response.get("step_duration"),
                }
            )
    return rows


def _archive_mobileuse_screens(step_list: Sequence[Any], screenshots_dir: Path) -> None:
    screenshots_dir.mkdir(parents=True, exist_ok=True)
    xml_dir = screenshots_dir.parent / "xml"
    image_counter = 1
    for index, step in enumerate(step_list, start=1):
        current_state = getattr(step, "curr_env_state", None)
        current_pixels = getattr(current_state, "pixels", None)
        if current_pixels is not None:
            stem = f"{image_counter:04d}_step_{index:02d}_before"
            current_pixels.save(screenshots_dir / f"{stem}.png")
            _write_env_state_dom_tree(current_state, xml_dir, f"{stem}.xml")
            image_counter += 1

        exec_state = getattr(step, "exec_env_state", None)
        exec_pixels = getattr(exec_state, "pixels", None)
        if exec_pixels is not None:
            stem = f"{image_counter:04d}_step_{index:02d}_after"
            exec_pixels.save(screenshots_dir / f"{stem}.png")
            _write_env_state_dom_tree(exec_state, xml_dir, f"{stem}.xml")
            image_counter += 1


def _copy_runtime_agent_dir(agent_dir: Path | None, report_dir: Path) -> dict[str, Any]:
    info = {"llm_log_rows": [], "usage_totals": aggregate_usage([])}
    if agent_dir is None or not agent_dir.exists():
        return info

    logs_dir = report_dir / "logs"
    screenshots_dir = report_dir / "screenshots"
    xml_dir = report_dir / "xml"
    raw_dir = report_dir / "raw"

    for log_name in ("appagent_log.jsonl", "mobileagentv2_log.jsonl"):
        log_path = agent_dir / log_name
        if not log_path.exists():
            continue
        rows = list(_iter_jsonl(log_path))
        _copy_file(log_path, raw_dir / log_name)
        _write_jsonl(logs_dir / "llm_interactions.jsonl", rows)
        info["llm_log_rows"] = rows
        info["usage_totals"] = _sum_usage_from_rows(rows)
        break

    _copy_tree_contents(agent_dir, screenshots_dir, patterns=["*.png", "*.jpg", "*.jpeg"])
    _copy_tree_contents(agent_dir / "screenshot", screenshots_dir, patterns=["*.png", "*.jpg", "*.jpeg"])
    _copy_tree_contents_preserve(agent_dir, screenshots_dir, patterns=["**/*.png", "**/*.jpg", "**/*.jpeg"])
    _copy_tree_contents(agent_dir, xml_dir, patterns=["*.xml"])
    _copy_tree_contents(agent_dir / "xml", xml_dir, patterns=["*.xml"])
    _copy_tree_contents_preserve(agent_dir, xml_dir, patterns=["**/*.xml"])
    _copy_tree_contents_preserve(
        agent_dir,
        raw_dir / "agent_log_dir",
        patterns=["**/*.jsonl", "**/*.json", "**/*.txt", "**/*.yaml", "**/*.yml", "**/*.md"],
    )
    return info


def _preserve_run_level_agent_artifacts(run_output_dir: Path) -> None:
    knowledge_artifacts = run_output_dir / "agent_logs" / "_knowledge_artifacts"
    if not knowledge_artifacts.exists():
        return

    knowledge_dir = run_output_dir / "knowledge"
    preserved_dir = knowledge_dir / "proactive_exploration"
    if preserved_dir.exists():
        shutil.rmtree(preserved_dir)
    knowledge_dir.mkdir(parents=True, exist_ok=True)
    shutil.copytree(knowledge_artifacts, preserved_dir)

    source_prefix = str(knowledge_artifacts)
    preserved_prefix = str(preserved_dir)

    def _rewrite_paths(value: Any) -> Any:
        if isinstance(value, str):
            return value.replace(source_prefix, preserved_prefix)
        if isinstance(value, list):
            return [_rewrite_paths(item) for item in value]
        if isinstance(value, dict):
            return {key: _rewrite_paths(item) for key, item in value.items()}
        return value

    for json_path in preserved_dir.rglob("*.json"):
        try:
            payload = json.loads(json_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        _write_json(json_path, _rewrite_paths(payload))

    manifest = preserved_dir / "proactive_exploration_manifest.json"
    if manifest.exists():
        shutil.copy2(manifest, knowledge_dir / "proactive_exploration_manifest.json")


def _load_androidworld_checkpoint(path: Path) -> dict[str, Any]:
    with gzip.open(path, "rb") as handle:
        payload = pickle.load(handle)
    if isinstance(payload, list) and payload:
        payload = payload[0]
    if not isinstance(payload, dict):
        raise ValueError(f"Unsupported checkpoint payload in {path}")
    return payload


def _androidworld_terminal_fact(payload: dict[str, Any]) -> tuple[str | None, str | None]:
    """Read a raw runner terminal event without interpreting task failure."""
    aux_data = payload.get("aux_data")
    if isinstance(aux_data, dict):
        terminal_event = aux_data.get("terminal_event")
        if terminal_event:
            return str(terminal_event), None
    if payload.get("exception_info"):
        return "exception", None
    return None, "terminal_event_not_recorded_by_runner"


def _androidworld_task_time_facts(payload: dict[str, Any]) -> dict[str, Any]:
    """Return task-boundary timestamps and explicit provenance from runner facts."""

    aux_data = payload.get("aux_data")
    if not isinstance(aux_data, dict):
        aux_data = {}
    started_at = aux_data.get("task_started_at")
    finished_at = aux_data.get("task_finished_at")
    finished_source = "runner_aux_data" if finished_at else None
    if not finished_at and payload.get("finish_dtime") is not None:
        finished_at = _to_jsonable(payload.get("finish_dtime"))
        finished_source = "checkpoint_finish_dtime"
    return {
        "task_started_at": str(started_at) if started_at else None,
        "task_started_at_source": "runner_aux_data" if started_at else None,
        "task_started_at_missing_reason": None if started_at else "task_started_at_not_recorded_by_runner",
        "task_finished_at": str(finished_at) if finished_at else None,
        "task_finished_at_source": finished_source,
        "task_finished_at_missing_reason": None if finished_at else "task_finished_at_not_recorded_by_runner",
    }


def _organize_androidworld_run(
    run_output_dir: Path,
    manifest: dict[str, Any],
    *,
    cleanup_raw_checkpoints: bool,
) -> list[dict[str, Any]]:
    run_name = run_output_dir.name
    strict_capture = bool(manifest.get("runtime_metrics_version"))
    report_root = run_output_dir / "task_reports"
    if report_root.exists():
        shutil.rmtree(report_root)
    report_root.mkdir(parents=True, exist_ok=True)

    checkpoint_paths = sorted(run_output_dir.glob("runs/**/*.pkl.gz"), key=lambda item: _natural_key(str(item)))
    agent_candidates = _discover_agent_dirs(run_output_dir / "agent_logs")
    records: list[dict[str, Any]] = []

    for checkpoint_path in checkpoint_paths:
        checkpoint_task_id = checkpoint_path.stem.replace(".pkl", "")
        notes: list[str] = []
        try:
            payload = _load_androidworld_checkpoint(checkpoint_path)
        except Exception as exc:
            payload = {
                "instance_id": checkpoint_task_id,
                "goal": checkpoint_task_id,
                "task_template": checkpoint_task_id,
                "episode_data": {},
                "is_successful": False,
                "exception_info": f"{type(exc).__name__}: {exc}",
            }
            notes.append("checkpoint_load_failed")
        # AndroidWorld stores a per-template numeric instance_id (often 0),
        # while the checkpoint filename carries the globally stable task id.
        task_id = checkpoint_task_id
        task = str(payload.get("goal") or payload.get("task_template") or task_id)
        report_dir = report_root / task_id
        report_dir.mkdir(parents=True, exist_ok=True)
        episode_data = payload.get("episode_data")
        if isinstance(episode_data, dict):
            step_list = list(episode_data.get("step_data") or [])
        else:
            step_list = []
            notes.append("episode_data_missing")

        _copy_terminal_log(run_output_dir, report_dir, notes)
        _copy_file(run_output_dir / "run_manifest.json", report_dir / "raw" / "run_manifest.json")
        _write_json(report_dir / "raw" / "checkpoint_summary.json", payload)
        _write_jsonl(report_dir / "logs" / "step_trace.jsonl", [_to_jsonable(step) for step in step_list])

        agent_dir = _agent_dir_from_step_artifacts(agent_candidates, step_list)
        if agent_dir is not None:
            agent_dir_source = "checkpoint_artifact_path"
        else:
            agent_dir = _match_agent_dir(agent_candidates, task=task, task_id=task_id)
            agent_dir_source = "task_text_fallback" if agent_dir is not None else "missing"
        runtime_info = _copy_runtime_agent_dir(agent_dir, report_dir)

        first_step = step_list[0] if step_list else None
        is_mobileuse = first_step is not None and hasattr(first_step, "vlm_call_history")
        execution_steps: list[str] = []
        usage_totals = runtime_info["usage_totals"]
        recovered_trace_rows: list[dict[str, Any]] = []
        recovered_mobileuse_steps: list[dict[str, Any]] = []
        effective_step_count = len(step_list)
        llm_ledger_source = "missing"

        if is_mobileuse:
            llm_rows = _build_mobileuse_llm_rows(step_list)
            llm_ledger_source = "checkpoint_mobileuse_step_data" if llm_rows else "missing"
            _write_jsonl(report_dir / "logs" / "llm_interactions.jsonl", llm_rows)
            _archive_mobileuse_screens(step_list, report_dir / "screenshots")
            usage_totals = _collect_mobileuse_usage_from_steps(step_list)
            execution_steps = [_step_summary_from_mobileuse(step, index) for index, step in enumerate(step_list, start=1)]
        else:
            llm_rows, llm_ledger_source = _select_runtime_llm_rows(
                step_list,
                runtime_info["llm_log_rows"],
            )
            if not step_list and not llm_rows:
                recovered_mobileuse_steps = _load_mobileuse_json_steps(agent_dir)
                if recovered_mobileuse_steps:
                    llm_rows = _build_mobileuse_llm_rows_from_json_steps(recovered_mobileuse_steps)
                    llm_ledger_source = "matched_agent_step_json_fallback"
            _write_jsonl(report_dir / "logs" / "llm_interactions.jsonl", llm_rows)
            if not _usage_has_known_tokens(usage_totals) and llm_rows:
                usage_totals = _sum_usage_from_rows(llm_rows)
            if not llm_rows:
                notes.append("llm_log_missing")
            if recovered_mobileuse_steps:
                _write_jsonl(report_dir / "logs" / "step_trace.jsonl", recovered_mobileuse_steps)
                execution_steps = [_step_summary_from_json_payload(step, index) for index, step in enumerate(recovered_mobileuse_steps, start=1)]
                effective_step_count = len(recovered_mobileuse_steps)
                notes.append("steps_recovered_from_mobileuse_agent_log")
            if not step_list and llm_rows:
                recovered_trace_rows = _trace_rows_from_llm_rows(llm_rows)
                if recovered_trace_rows and not recovered_mobileuse_steps:
                    _write_jsonl(report_dir / "logs" / "step_trace.jsonl", recovered_trace_rows)
                    execution_steps = [_step_summary_from_trace(line, index) for index, line in enumerate(recovered_trace_rows, start=1)]
                    effective_step_count = len(recovered_trace_rows)
                    notes.append("steps_recovered_from_runtime_log")
            if not execution_steps:
                execution_steps = [_step_summary_from_runtime_step(step, index) for index, step in enumerate(step_list, start=1)]

        _write_androidworld_step_dirs(
            report_dir,
            step_list,
            llm_rows,
            strict_capture=strict_capture,
        )
        if recovered_trace_rows:
            _write_generic_step_dirs(
                report_dir,
                recovered_trace_rows,
                llm_rows,
                strict_capture=strict_capture,
            )
        if not step_list and recovered_mobileuse_steps:
            _write_generic_step_dirs(
                report_dir,
                recovered_mobileuse_steps,
                llm_rows,
                strict_capture=strict_capture,
            )

        if not _usage_has_known_tokens(usage_totals):
            notes.append("token_usage_missing")
        usage_by_stage = _sum_usage_by_stage_from_rows(llm_rows)

        # The canonical cost source for new reports is the call-level ledger.
        # Step and run totals are derived from it, never added on top of it.
        canonical_usage = aggregate_usage(llm_rows)
        if canonical_usage.get("call_count", 0):
            usage_totals = canonical_usage
        step_usage_rows: list[dict[str, Any]] = []
        grouped_calls = _llm_rows_by_step(llm_rows)
        max_step_for_ledger = max([*grouped_calls.keys(), effective_step_count, 0])
        for step_index in range(1, max_step_for_ledger + 1):
            step_usage = aggregate_usage(grouped_calls.get(step_index, []))
            step_usage["step"] = step_index
            step_usage_rows.append(step_usage)
        _write_jsonl(report_dir / "logs" / "step_usage.jsonl", step_usage_rows)
        token_reconciliation = reconcile_token_layers(llm_rows, step_usage_rows, usage_totals)

        if not list((report_dir / "screenshots").glob("*")):
            notes.append("screenshots_partial_or_missing")

        runtime_seconds = payload.get("run_time")
        if not isinstance(runtime_seconds, (int, float)):
            durations = [_extract_step_duration(step) for step in step_list]
            if not durations and recovered_trace_rows:
                durations = [_extract_step_duration(step) for step in recovered_trace_rows]
            durations = [value for value in durations if value is not None]
            runtime_seconds = sum(durations) if durations else None
        step_durations = [
            _extract_step_duration(step)
            for step in (step_list or recovered_trace_rows)
        ]
        step_durations = [value for value in step_durations if value is not None]
        sum_step_duration_seconds = sum(step_durations) if step_durations else None
        runtime_delta_seconds = (
            float(runtime_seconds) - sum_step_duration_seconds
            if runtime_seconds is not None and sum_step_duration_seconds is not None
            else None
        )
        if runtime_delta_seconds is None:
            runtime_reason = "step_or_run_duration_missing"
        elif abs(runtime_delta_seconds) <= 0.05:
            runtime_reason = "within_tolerance"
        elif runtime_delta_seconds < 0:
            runtime_reason = "sum_step_duration_exceeds_run_runtime"
        else:
            runtime_reason = "run_runtime_includes_task_initialization_and_evaluation_outside_step_durations"
        runtime_reconciliation = {
            "status": "ok" if runtime_delta_seconds is not None and abs(runtime_delta_seconds) <= 0.05 else (
                "difference_recorded" if runtime_delta_seconds is not None else "not_checkable"
            ),
            "run_runtime_seconds": runtime_seconds,
            "sum_step_duration_seconds": sum_step_duration_seconds,
            "delta_run_minus_steps_seconds": runtime_delta_seconds,
            "non_step_runtime_seconds": runtime_delta_seconds,
            "reason": runtime_reason,
        }

        success = _coerce_success(payload.get("is_successful"))
        final_result = "success" if success else "failed"
        terminal_event, terminal_event_missing_reason = _androidworld_terminal_fact(payload)
        task_time_facts = _androidworld_task_time_facts(payload)
        aux_data = payload.get("aux_data") if isinstance(payload.get("aux_data"), dict) else {}
        started_at = task_time_facts["task_started_at"]
        finished_at = task_time_facts["task_finished_at"]

        summary = {
            "task_id": task_id,
            "task": task,
            "checkpoint_path": str(checkpoint_path),
            "source_checkpoint_path": str(checkpoint_path),
            "checkpoint_path_exists_at_report_time": checkpoint_path.exists(),
            "checkpoint_retained_after_reporting": not cleanup_raw_checkpoints,
            "agent_log_dir": str(agent_dir) if agent_dir else None,
            "agent_log_dir_source": agent_dir_source,
            "agent_log_dir_exists_at_report_time": bool(agent_dir and agent_dir.exists()),
            "llm_ledger_source": llm_ledger_source,
            "success": success,
            "evaluator_result": payload.get("is_successful"),
            "terminal_event": terminal_event,
            "terminal_event_missing_reason": terminal_event_missing_reason,
            "evaluation_state_facts": aux_data.get("evaluation_state_facts"),
            "runtime_seconds": runtime_seconds,
            **task_time_facts,
            "step_count": effective_step_count,
            "usage": usage_totals,
            "usage_by_stage": usage_by_stage,
            "token_reconciliation": token_reconciliation,
            "configured_max_steps": manifest.get("configured_max_steps"),
            "runtime_reconciliation": runtime_reconciliation,
            "step_durations_seconds": step_durations,
            "execution_steps": execution_steps,
        }
        _write_json(report_dir / "summary.json", summary)

        records.append(
            _make_base_record(
                manifest=manifest,
                benchmark="androidworld",
                mode=str(manifest.get("mode") or run_output_dir.parent.name),
                run_name=run_name,
                run_output_dir=run_output_dir,
                task_id=task_id,
                task=task,
                execution_steps="\n".join(execution_steps),
                success=success,
                final_result=final_result,
                step_count=effective_step_count,
                usage_totals=usage_totals,
                usage_by_stage=usage_by_stage,
                runtime_seconds=runtime_seconds,
                started_at=started_at,
                finished_at=finished_at,
                report_dir=report_dir,
                notes=notes,
                token_reconciliation=token_reconciliation,
                runtime_reconciliation=runtime_reconciliation,
                terminal_event=terminal_event,
                terminal_event_missing_reason=terminal_event_missing_reason,
                started_at_source=task_time_facts["task_started_at_source"],
                started_at_missing_reason=task_time_facts["task_started_at_missing_reason"],
                finished_at_source=task_time_facts["task_finished_at_source"],
                finished_at_missing_reason=task_time_facts["task_finished_at_missing_reason"],
            )
        )

        if cleanup_raw_checkpoints and checkpoint_path.exists():
            checkpoint_path.unlink()

    for runs_dir in sorted(run_output_dir.glob("runs/*"), key=lambda item: _natural_key(item.name), reverse=True):
        if runs_dir.is_dir() and not any(runs_dir.iterdir()):
            runs_dir.rmdir()

    return records


def _extract_androidlab_task_id(task_dir: Path) -> str:
    return re.sub(r"_\d{4}-\d{2}-\d{2}_\d{2}-\d{2}-\d{2}$", "", task_dir.name)


def _find_androidlab_task_dirs(run_output_dir: Path, run_name: str) -> list[Path]:
    evaluation_root = run_output_dir / "evaluation" / run_name
    if not evaluation_root.exists():
        return []
    return sorted([path for path in evaluation_root.iterdir() if path.is_dir()], key=lambda item: _natural_key(item.name))


def _find_androidlab_agent_dir(run_output_dir: Path, task_dir: Path, task: str, task_id: str) -> Path | None:
    direct_candidates = [path for path in task_dir.iterdir() if path.is_dir() and path.name not in {"Screen", "traces", "xml"}]
    if direct_candidates:
        return direct_candidates[0]
    shared_candidates = _discover_agent_dirs(
        run_output_dir / "evaluation" / "agent_logs",
        run_output_dir / "agent_logs",
    )
    return _match_agent_dir(shared_candidates, task=task, task_id=task_id)


def _organize_androidlab_run(run_output_dir: Path, manifest: dict[str, Any]) -> list[dict[str, Any]]:
    run_name = str(manifest.get("actual_run_name") or run_output_dir.name)
    report_root = run_output_dir / "task_reports"
    if report_root.exists():
        shutil.rmtree(report_root)
    report_root.mkdir(parents=True, exist_ok=True)

    task_catalog = _get_androidlab_task_catalog()
    task_dirs = _find_androidlab_task_dirs(run_output_dir, run_name)
    records: list[dict[str, Any]] = []

    for task_dir in task_dirs:
        trace_path = task_dir / "traces" / "trace.jsonl"
        trace_rows = list(_iter_jsonl(trace_path))
        task_id = _extract_androidlab_task_id(task_dir)
        catalog_entry = task_catalog.get(task_id) or {}
        task = str(catalog_entry.get("task") or (trace_rows[0].get("target") if trace_rows else task_id))
        report_dir = report_root / task_id
        report_dir.mkdir(parents=True, exist_ok=True)
        notes: list[str] = []

        _copy_terminal_log(run_output_dir, report_dir, notes)
        _copy_file(run_output_dir / "run_manifest.json", report_dir / "raw" / "run_manifest.json")
        _copy_file(trace_path, report_dir / "raw" / "trace.jsonl")
        _write_jsonl(report_dir / "logs" / "step_trace.jsonl", trace_rows)
        _copy_tree_contents(task_dir / "Screen", report_dir / "screenshots", patterns=["*.png", "*.jpg", "*.jpeg"])
        _copy_tree_contents(task_dir / "xml", report_dir / "xml", patterns=["*.xml"])
        if not list((report_dir / "screenshots").glob("*")):
            notes.append("screenshots_missing")

        agent_dir = _find_androidlab_agent_dir(run_output_dir, task_dir, task, task_id)
        runtime_info = _copy_runtime_agent_dir(agent_dir, report_dir)

        llm_rows = runtime_info["llm_log_rows"] or _build_trace_llm_rows(trace_rows)
        if runtime_info["llm_log_rows"]:
            _copy_file(next((agent_dir / name for name in ("appagent_log.jsonl", "mobileagentv2_log.jsonl") if (agent_dir / name).exists()), None), report_dir / "raw" / "agent_log.jsonl")
        else:
            notes.append("runtime_agent_log_missing")
        _write_jsonl(report_dir / "logs" / "llm_interactions.jsonl", llm_rows)
        _write_androidlab_step_dirs(report_dir, trace_rows, llm_rows)

        usage_totals = _sum_usage_from_rows(trace_rows)
        if not _usage_has_known_tokens(usage_totals):
            if llm_rows:
                usage_totals = _sum_usage_from_rows(llm_rows)
        if not _usage_has_known_tokens(usage_totals):
            usage_totals = runtime_info["usage_totals"]
            if not _usage_has_known_tokens(usage_totals):
                notes.append("token_usage_missing")

        screenshot_epochs = [_parse_screenshot_epoch(path) for path in sorted((task_dir / "Screen").glob("*"), key=lambda item: _natural_key(item.name))]
        screenshot_epochs = [value for value in screenshot_epochs if value is not None]
        runtime_seconds = None
        if len(screenshot_epochs) >= 2:
            runtime_seconds = screenshot_epochs[-1] - screenshot_epochs[0]

        success = _evaluate_androidlab_task(task_id, task_dir, task_catalog)
        last_response = trace_rows[-1].get("current_response", {}) if trace_rows else {}
        agent_finished = bool(last_response.get("done")) if isinstance(last_response, dict) else False
        if success:
            final_result = "success"
        elif agent_finished:
            final_result = "agent_finished_but_metric_failed"
        else:
            final_result = "incomplete"

        execution_steps = [_step_summary_from_trace(line, index) for index, line in enumerate(trace_rows, start=1)]

        usage_by_stage = _sum_usage_by_stage_from_rows(llm_rows or trace_rows)
        summary = {
            "task_id": task_id,
            "task": task,
            "task_dir": str(task_dir),
            "agent_log_dir": str(agent_dir) if agent_dir else None,
            "success": success,
            "final_result": final_result,
            "runtime_seconds": runtime_seconds,
            "step_count": len(trace_rows),
            "usage": usage_totals,
            "usage_by_stage": usage_by_stage,
            "execution_steps": execution_steps,
        }
        _write_json(report_dir / "summary.json", summary)

        records.append(
            _make_base_record(
                manifest=manifest,
                benchmark="androidlab",
                mode=str(manifest.get("mode") or run_output_dir.parent.name),
                run_name=run_name,
                run_output_dir=run_output_dir,
                task_id=task_id,
                task=task,
                execution_steps="\n".join(execution_steps),
                success=success,
                final_result=final_result,
                step_count=len(trace_rows),
                usage_totals=usage_totals,
                usage_by_stage=usage_by_stage,
                runtime_seconds=runtime_seconds,
                started_at=manifest.get("started_at"),
                finished_at=manifest.get("finished_at"),
                report_dir=report_dir,
                notes=notes,
            )
        )

    return records


def organize_run_output(
    run_output_dir: str | Path,
    *,
    append_excel: bool = True,
    cleanup_raw_checkpoints: bool = True,
) -> list[dict[str, Any]]:
    run_output_dir = Path(run_output_dir).resolve()
    manifest = _load_json(run_output_dir / "run_manifest.json") or {}
    benchmark = str(manifest.get("benchmark") or run_output_dir.parent.parent.name)
    benchmark_aliases = {
        "androidworld": "androidworld",
        "安卓世界": "androidworld",
        "androidlab": "androidlab",
        "android-lab": "androidlab",
        "安卓实验室": "androidlab",
    }
    benchmark = benchmark_aliases.get(benchmark.lower(), benchmark_aliases.get(benchmark, benchmark))
    existing_run_report = _load_json(run_output_dir / "run_report.json") or {}
    if existing_run_report.get("records"):
        has_androidworld_raw = any(run_output_dir.glob("runs/**/*.pkl.gz"))
        has_androidlab_raw = (run_output_dir / "evaluation").exists()
        if (benchmark == "androidworld" and not has_androidworld_raw) or (benchmark == "androidlab" and not has_androidlab_raw):
            records = list(existing_run_report["records"])
            _backfill_step_dirs_from_task_reports(run_output_dir, benchmark)
            _preserve_run_level_agent_artifacts(run_output_dir)
            compact_run_output(run_output_dir)
            if _should_localize_large_run_files():
                localize_run_output_files(run_output_dir)
            if append_excel and records:
                outputs_root = _find_outputs_root(run_output_dir)
                append_records_to_benchmark_outputs(outputs_root, benchmark, records)
            return records

    if benchmark == "androidworld":
        records = _organize_androidworld_run(
            run_output_dir,
            manifest,
            cleanup_raw_checkpoints=cleanup_raw_checkpoints,
        )
    elif benchmark == "androidlab":
        records = _organize_androidlab_run(run_output_dir, manifest)
    else:
        raise ValueError(f"Unsupported benchmark={benchmark!r} for {run_output_dir}")

    run_report = {
        "run_output_dir": str(run_output_dir),
        "benchmark": benchmark,
        "mode": manifest.get("mode"),
        "run_name": run_output_dir.name,
        "records": records,
    }
    _write_json(run_output_dir / "run_report.json", run_report)
    _preserve_run_level_agent_artifacts(run_output_dir)
    compact_run_output(run_output_dir)
    if _should_localize_large_run_files():
        localize_run_output_files(run_output_dir)

    if append_excel and records:
        outputs_root = _find_outputs_root(run_output_dir)
        append_records_to_benchmark_outputs(outputs_root, benchmark, records)
    return records


def scan_run_output_dirs(outputs_root: str | Path) -> list[Path]:
    outputs_root = Path(outputs_root).resolve()
    run_dirs: list[Path] = []
    for benchmark in ("androidworld", "androidlab"):
        benchmark_root = outputs_root / benchmark
        if not benchmark_root.exists():
            continue
        for mode_dir in sorted([path for path in benchmark_root.iterdir() if path.is_dir()], key=lambda item: _natural_key(item.name)):
            for run_dir in sorted([path for path in mode_dir.iterdir() if path.is_dir()], key=lambda item: _natural_key(item.name)):
                if (run_dir / "run_manifest.json").exists():
                    run_dirs.append(run_dir)
    return run_dirs


def rebuild_reports(
    outputs_root: str | Path,
    *,
    keep_run_prefixes: Sequence[str] | None = None,
) -> list[dict[str, Any]]:
    outputs_root = Path(outputs_root).resolve()
    translator = ChineseTranslator()
    kept_records: list[dict[str, Any]] = []
    records_by_benchmark: dict[str, list[dict[str, Any]]] = {"androidworld": [], "androidlab": []}
    for run_dir in scan_run_output_dirs(outputs_root):
        if keep_run_prefixes and not any(run_dir.name.startswith(prefix) for prefix in keep_run_prefixes):
            continue
        run_records = organize_run_output(
            run_dir,
            append_excel=False,
            cleanup_raw_checkpoints=True,
        )
        kept_records.extend(run_records)
        for record in run_records:
            benchmark = str(record.get("benchmark") or "").lower()
            if benchmark in records_by_benchmark:
                records_by_benchmark[benchmark].append(record)

    for benchmark, records in records_by_benchmark.items():
        paths = benchmark_report_paths(outputs_root, benchmark)
        write_records_to_excel(paths["excel"], records)
        localized_records = [localize_record(record, translator) for record in records]
        _write_json(paths["index"], localized_records)
    if _should_localize_large_run_files():
        localize_outputs_tree(outputs_root)
    return kept_records


def prune_outputs(outputs_root: str | Path, *, keep_run_prefixes: Sequence[str]) -> list[str]:
    outputs_root = Path(outputs_root).resolve()
    removed: list[str] = []

    for benchmark in ("androidworld", "androidlab"):
        benchmark_root = outputs_root / benchmark
        if not benchmark_root.exists():
            continue
        for mode_dir in [path for path in benchmark_root.iterdir() if path.is_dir()]:
            for child in list(mode_dir.iterdir()):
                if child.is_dir():
                    if any(child.name.startswith(prefix) for prefix in keep_run_prefixes):
                        continue
                    shutil.rmtree(child)
                    removed.append(str(child))
                else:
                    child.unlink()
                    removed.append(str(child))

    for extra_name in ("androidlab_avd_logs", "emulator_logs"):
        extra_path = outputs_root / extra_name
        if extra_path.exists():
            shutil.rmtree(extra_path)
            removed.append(str(extra_path))

    for child in list(outputs_root.iterdir()):
        if child.name in {
            "androidworld",
            "androidlab",
            "androidworld_results.xlsx",
            "androidlab_results.xlsx",
            "androidworld_report_index.json",
            "androidlab_report_index.json",
        }:
            continue
        if child.name.endswith("_analysis.md"):
            continue
        if child.is_file():
            child.unlink()
            removed.append(str(child))

    return removed


def compact_run_output(run_output_dir: str | Path) -> list[str]:
    run_output_dir = Path(run_output_dir).resolve()
    manifest = _load_json(run_output_dir / "run_manifest.json") or {}
    benchmark = str(manifest.get("benchmark") or run_output_dir.parent.parent.name)
    removed: list[str] = []

    removable_dirs = [run_output_dir / "agent_logs"]
    if benchmark == "androidworld":
        removable_dirs.append(run_output_dir / "runs")
    if benchmark == "androidlab":
        removable_dirs.append(run_output_dir / "evaluation")

    for path in removable_dirs:
        if path.exists():
            shutil.rmtree(path)
            removed.append(str(path))

    return removed
